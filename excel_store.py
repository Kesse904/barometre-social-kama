"""Lecture du questionnaire et écriture des réponses dans le classeur Excel.

L'écriture se fait directement dans le XML du fichier .xlsx (et non via openpyxl)
afin de conserver intacts le graphique, les commentaires et la mise en forme.
"""
import os
import re
import tempfile
import threading
import zipfile
from xml.sax.saxutils import escape

import openpyxl

SHEET_QUESTIONNAIRE = "Questionnaire"
SHEET_SAISIE = "Saisie des réponses"
SHEET_RESULTATS = "Résultats"

FIRST_DATA_ROW = 2
FORMULA_LAST_ROW = 1000  # les formules de "Résultats" couvrent jusqu'à cette ligne

COL_ID, COL_SERVICE, COL_ANCIENNETE = "A", "B", "C"
COL_COMMENTAIRE = "AB"
NOTE_STYLE = "11"  # style des cellules jaunes de saisie

_lock = threading.Lock()


def col_to_index(col: str) -> int:
    n = 0
    for ch in col:
        n = n * 26 + (ord(ch) - 64)
    return n


def index_to_col(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def read_questionnaire(path: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[SHEET_QUESTIONNAIRE]
    questions = []
    for num, theme, texte, *_ in ws.iter_rows(min_row=5, values_only=True):
        if isinstance(num, (int, float)) and texte:
            questions.append({"numero": int(num), "theme": theme, "texte": texte})
    wb.close()
    return questions


def _sheet_paths(z: zipfile.ZipFile) -> dict[str, str]:
    """Nom de feuille -> chemin XML dans l'archive."""
    wb_xml = z.read("xl/workbook.xml").decode("utf-8")
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    targets = dict(re.findall(r'<Relationship Id="([^"]+)"[^>]*?Target="([^"]+)"', rels))
    targets.update({k: v for v, k in re.findall(r'Target="([^"]+)"[^>]*?Id="([^"]+)"', rels)})
    result = {}
    for name, rid in re.findall(r'<sheet name="([^"]+)"[^>]*?r:id="([^"]+)"', wb_xml):
        name = name.replace("&amp;", "&")
        result[name] = "xl/" + targets[rid].lstrip("/").removeprefix("xl/")
    return result


_CELL_RE = re.compile(r'<c r="([A-Z]+)(\d+)"[^>]*?(?:/>|>.*?</c>)', re.S)


def _parse_cells(row_body: str) -> dict[str, str]:
    return {m.group(1): m.group(0) for m in _CELL_RE.finditer(row_body)}


def _cell_has_value(xml: str) -> bool:
    return "<v>" in xml or "<is>" in xml


def _num_cell(ref: str, value, style: str | None = None) -> str:
    s = f' s="{style}"' if style else ""
    return f'<c r="{ref}"{s}><v>{value}</v></c>'


def _str_cell(ref: str, value: str) -> str:
    return f'<c r="{ref}" t="inlineStr"><is><t xml:space="preserve">{escape(value)}</t></is></c>'


def _write_row(sheet_xml: str, reponse: dict, nb_questions: int) -> tuple[str, int]:
    rows = {int(m.group(1)): m for m in re.finditer(r'<row r="(\d+)"[^>]*?(?:/>|>(.*?)</row>)', sheet_xml, re.S)}
    first_q, last_q = col_to_index("D"), col_to_index("D") + nb_questions - 1

    # Première ligne dont aucune note n'est encore saisie
    target = FIRST_DATA_ROW
    while True:
        m = rows.get(target)
        if m is None:
            break
        cells = _parse_cells(m.group(2) or "")
        if not any(_cell_has_value(cells.get(index_to_col(i), "")) for i in range(first_q, last_q + 1)):
            break
        target += 1

    existing = rows.get(target)
    cells = _parse_cells(existing.group(2) or "") if existing else {}
    if COL_ID not in cells or not _cell_has_value(cells[COL_ID]):
        cells[COL_ID] = _num_cell(f"{COL_ID}{target}", target - 1)
    if reponse.get("service"):
        cells[COL_SERVICE] = _str_cell(f"{COL_SERVICE}{target}", reponse["service"])
    if reponse.get("anciennete") is not None:
        cells[COL_ANCIENNETE] = _num_cell(f"{COL_ANCIENNETE}{target}", reponse["anciennete"])
    for i, note in enumerate(reponse["notes"]):
        col = index_to_col(first_q + i)
        cells[col] = _num_cell(f"{col}{target}", int(note), NOTE_STYLE)
    if reponse.get("commentaire"):
        cells[COL_COMMENTAIRE] = _str_cell(f"{COL_COMMENTAIRE}{target}", reponse["commentaire"])

    body = "".join(cells[c] for c in sorted(cells, key=col_to_index))
    new_row = f'<row r="{target}">{body}</row>'

    if existing:
        sheet_xml = sheet_xml[: existing.start()] + new_row + sheet_xml[existing.end() :]
    else:
        after = [m for r, m in rows.items() if r > target]
        if after:
            pos = min(after, key=lambda m: m.start()).start()
        else:
            pos = sheet_xml.index("</sheetData>")
        sheet_xml = sheet_xml[:pos] + new_row + sheet_xml[pos:]
        last_row = max(max(rows, default=1), target)
        sheet_xml = re.sub(r'<dimension ref="A1:([A-Z]+)\d+"/>', rf'<dimension ref="A1:\g<1>{last_row}"/>', sheet_xml, count=1)
    return sheet_xml, target


def _fix_result_formulas(xml: str) -> str:
    """Les formules d'origine commencent à la ligne 3 (le 1er répondant était ignoré)
    et s'arrêtent à la ligne 41 : on les étend à 2:FORMULA_LAST_ROW."""
    def repl(m):
        return f"'{SHEET_SAISIE}'!{m.group(1)}{FIRST_DATA_ROW}:{m.group(2)}{FORMULA_LAST_ROW}"
    return re.sub(r"'Saisie des réponses'!([A-Z]+)\d+:([A-Z]+)\d+", repl, xml)


def _force_recalc(wb_xml: str) -> str:
    if "fullCalcOnLoad" in wb_xml:
        return wb_xml
    return re.sub(r"<calcPr\b", '<calcPr fullCalcOnLoad="1"', wb_xml, count=1)


def append_reponses(path: str, reponses: list[dict], nb_questions: int) -> list[int]:
    """Ajoute des réponses dans la feuille de saisie. Lève PermissionError si le fichier est verrouillé."""
    with _lock:
        with zipfile.ZipFile(path) as z:
            paths = _sheet_paths(z)
            saisie, resultats = paths[SHEET_SAISIE], paths[SHEET_RESULTATS]
            entries = {info.filename: (info, z.read(info.filename)) for info in z.infolist()}

        sheet_xml = entries[saisie][1].decode("utf-8")
        lignes = []
        for rep in reponses:
            sheet_xml, ligne = _write_row(sheet_xml, rep, nb_questions)
            lignes.append(ligne)

        modified = {
            saisie: sheet_xml.encode("utf-8"),
            resultats: _fix_result_formulas(entries[resultats][1].decode("utf-8")).encode("utf-8"),
            "xl/workbook.xml": _force_recalc(entries["xl/workbook.xml"][1].decode("utf-8")).encode("utf-8"),
        }

        fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=os.path.dirname(os.path.abspath(path)))
        os.close(fd)
        try:
            with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as out:
                for name, (info, data) in entries.items():
                    out.writestr(info, modified.get(name, data))
            # Échoue avec PermissionError si le classeur est ouvert dans Excel
            os.replace(tmp, path)
        finally:
            if os.path.exists(tmp):
                os.remove(tmp)
        return lignes


def lecture(score: float | None) -> str:
    if score is None:
        return ""
    if score < 2.5:
        return "Alerte"
    if score < 3.5:
        return "Fragile"
    if score < 4.2:
        return "Correct"
    return "Solide"


def read_reponses(path: str, nb_questions: int) -> list[dict]:
    """Relit les réponses saisies dans la feuille de saisie."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[SHEET_SAISIE]
    lignes = []
    for num_ligne, row in enumerate(ws.iter_rows(min_row=FIRST_DATA_ROW, values_only=True), start=FIRST_DATA_ROW):
        row = list(row) + [None] * (28 - len(row))
        notes = row[3 : 3 + nb_questions]
        if not any(isinstance(n, (int, float)) for n in notes):
            continue
        lignes.append({
            "ligne": num_ligne,
            "id": row[0],
            "service": row[1],
            "anciennete": row[2],
            "notes": notes,
            "commentaire": row[col_to_index(COL_COMMENTAIRE) - 1],
        })
    wb.close()
    return lignes


def calcul_resultats(questions: list[dict], lignes: list[dict]) -> dict:
    """Calcule les scores comme l'onglet "Résultats" : score d'un thème = moyenne de
    toutes ses notes ; score global = moyenne des thèmes."""
    resultat_themes = []
    for q_theme in dict.fromkeys(q["theme"] for q in questions):
        idx = [q["numero"] - 1 for q in questions if q["theme"] == q_theme]
        valeurs = [l["notes"][i] for l in lignes for i in idx if isinstance(l["notes"][i], (int, float))]
        score = sum(valeurs) / len(valeurs) if valeurs else None
        resultat_themes.append({"theme": q_theme, "score": score, "lecture": lecture(score)})

    scores = [t["score"] for t in resultat_themes if t["score"] is not None]
    global_score = sum(scores) / len(scores) if scores else None
    return {
        "repondants": len(lignes),
        "themes": resultat_themes,
        "global": {"score": global_score, "lecture": lecture(global_score)},
        "reponses": lignes,
    }
